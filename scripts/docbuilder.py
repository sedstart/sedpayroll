"""
Shared helpers for generating the SedPayroll requirements documents
(BRD/PRD/SRS) as formatted .docx files with python-docx.

Not a general-purpose library — just enough structure (cover page, styled
headings, requirement tables, a native Word TOC field) to keep
generate_requirements_docs.py focused on content instead of formatting
boilerplate.
"""
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.section import WD_SECTION

BRAND_INDIGO = RGBColor(0x43, 0x38, 0xCA)  # indigo-700-ish
BRAND_INDIGO_DARK = RGBColor(0x31, 0x2E, 0x81)
BRAND_GRAY = RGBColor(0x37, 0x41, 0x51)
BRAND_MUTED = RGBColor(0x6B, 0x72, 0x80)
HEADER_FILL = "4338CA"
ALT_ROW_FILL = "EEF2FF"


def new_document():
    doc = Document()

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BRAND_GRAY
    normal.paragraph_format.space_after = Pt(6)

    for level, size, color in [
        (1, 20, BRAND_INDIGO_DARK),
        (2, 15, BRAND_INDIGO_DARK),
        (3, 12.5, BRAND_INDIGO),
        (4, 11, BRAND_GRAY),
    ]:
        h = styles[f"Heading {level}"]
        h.font.name = "Calibri"
        h.font.size = Pt(size)
        h.font.color.rgb = color
        h.font.bold = True
        h.paragraph_format.space_before = Pt(18 if level == 1 else 12)
        h.paragraph_format.space_after = Pt(8 if level <= 2 else 4)
        h.paragraph_format.keep_with_next = True

    section = doc.sections[0]
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)

    return doc


def _field(paragraph, instr_text, placeholder_text):
    run = paragraph.add_run()
    r = run._r

    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instr_text

    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")

    text_el = OxmlElement("w:t")
    text_el.text = placeholder_text

    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")

    r.append(fld_begin)
    r.append(instr)
    r.append(fld_sep)
    r.append(text_el)
    r.append(fld_end)


def add_toc(doc):
    p = doc.add_paragraph()
    _field(
        p,
        'TOC \\o "1-3" \\h \\z \\u',
        "Table of contents — right-click and choose “Update Field” "
        "(or Update Table) in Word to generate this list.",
    )


def add_page_number_footer(doc):
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Page ").font.size = Pt(9)
    _field(p, "PAGE", "1")
    p.add_run(" of ").font.size = Pt(9)
    _field(p, "NUMPAGES", "1")
    for run in p.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = BRAND_MUTED


def add_cover_page(doc, *, doc_type, title, subtitle, version, date_str, status):
    for _ in range(4):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("SedPayroll")
    run.font.size = Pt(16)
    run.font.color.rgb = BRAND_MUTED
    run.font.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(doc_type)
    run.font.size = Pt(34)
    run.font.color.rgb = BRAND_INDIGO_DARK
    run.font.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.font.size = Pt(16)
    run.font.color.rgb = BRAND_GRAY

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(subtitle)
    run.font.size = Pt(11)
    run.font.color.rgb = BRAND_MUTED
    run.font.italic = True

    for _ in range(6):
        doc.add_paragraph()

    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    rows = [
        ("Version", version),
        ("Date", date_str),
        ("Status", status),
        ("Product", "SedPayroll — payroll, attendance & employee management"),
        ("Prepared by", "Product & Engineering"),
    ]
    for label, value in rows:
        row = table.add_row()
        row.cells[0].width = Cm(4)
        row.cells[1].width = Cm(9)
        run = row.cells[0].paragraphs[0].add_run(label)
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = BRAND_GRAY
        run2 = row.cells[1].paragraphs[0].add_run(value)
        run2.font.size = Pt(10)
        run2.font.color.rgb = BRAND_GRAY

    doc.add_page_break()


def add_document_control(doc, *, history_rows, reviewers_rows=None):
    doc.add_heading("Document Control", level=2)
    doc.add_heading("Revision history", level=3)
    add_table(doc, ["Version", "Date", "Author", "Description"], history_rows)

    if reviewers_rows:
        doc.add_heading("Reviewers / approvers", level=3)
        add_table(doc, ["Name / Role", "Area", "Status"], reviewers_rows)

    doc.add_page_break()


def _shade_cell(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tcPr.append(shd)


def add_table(doc, headers, rows, col_widths=None, header_fill=HEADER_FILL):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True

    hdr_cells = table.rows[0].cells
    for i, text in enumerate(headers):
        hdr_cells[i].text = ""
        run = hdr_cells[i].paragraphs[0].add_run(text)
        run.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _shade_cell(hdr_cells[i], header_fill)

    for r_i, row_data in enumerate(rows):
        row_cells = table.add_row().cells
        for i, value in enumerate(row_data):
            cell = row_cells[i]
            cell.text = ""
            lines = str(value).split("\n")
            p = cell.paragraphs[0]
            for li, line in enumerate(lines):
                if li > 0:
                    p = cell.add_paragraph()
                run = p.add_run(line)
                run.font.size = Pt(9.5)
                run.font.color.rgb = BRAND_GRAY
            if r_i % 2 == 1:
                _shade_cell(cell, ALT_ROW_FILL)

    if col_widths:
        for i, w in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = Cm(w)

    doc.add_paragraph()
    return table


def add_bullets(doc, items, style="List Bullet"):
    for item in items:
        doc.add_paragraph(item, style=style)


def add_numbered(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Number")
