#!/usr/bin/env python3
"""
Generates the three SedPayroll requirements documents as .docx files:
  docs/SedPayroll-BRD.docx  (Business Requirements Document)
  docs/SedPayroll-PRD.docx  (Product Requirements Document)
  docs/SedPayroll-SRS.docx  (Software Requirements Specification)

Content reflects the CURRENT, already-built implementation (reverse
documented from the source under app/, lib/, db/schema.ts as of the date
below) rather than a green-field proposal — each doc says so up front.

Run: python3 scripts/generate_requirements_docs.py
Requires: pip install python-docx
"""
import subprocess
import sys
from pathlib import Path

try:
    import docx  # noqa: F401
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", "python-docx"])

sys.path.insert(0, str(Path(__file__).resolve().parent))
from docbuilder import (  # noqa: E402
    new_document,
    add_cover_page,
    add_document_control,
    add_toc,
    add_table,
    add_bullets,
    add_numbered,
    add_page_number_footer,
)

ROOT = Path(__file__).resolve().parent.parent
DOCS_DIR = ROOT / "docs"
VERSION = "1.0"
DATE_STR = "05 September 2026"
STATUS = "Draft — for review"

HISTORY_ROWS = [
    (VERSION, DATE_STR, "Product & Engineering",
     "Initial version, documenting the system as currently implemented."),
]


def req_table(doc, headers, rows, col_widths=None):
    add_table(doc, headers, rows, col_widths=col_widths)


# ---------------------------------------------------------------------------
# BRD
# ---------------------------------------------------------------------------

def build_brd():
    doc = new_document()
    add_cover_page(
        doc,
        doc_type="BRD",
        title="Business Requirements Document",
        subtitle="Payroll, attendance and employee management — current-state business requirements",
        version=VERSION,
        date_str=DATE_STR,
        status=STATUS,
    )
    add_document_control(doc, history_rows=HISTORY_ROWS)
    add_toc(doc)
    doc.add_page_break()

    # 1. Introduction
    doc.add_heading("1. Introduction", level=1)
    doc.add_heading("1.1 Purpose", level=2)
    doc.add_paragraph(
        "This Business Requirements Document (BRD) records the business needs that "
        "SedPayroll — an internal payroll, attendance and employee-management portal — "
        "exists to satisfy. It is written against the system as already built, so it "
        "doubles as a current-state baseline: a shared reference for the business "
        "before any further change is scoped, and an onboarding reference for anyone "
        "new to the product."
    )
    doc.add_heading("1.2 Business background", level=2)
    doc.add_paragraph(
        "Organizations running payroll and attendance across spreadsheets and ad-hoc "
        "tools face slow, error-prone monthly payroll cycles, no self-service for "
        "employees, and no single source of truth for who is active, who is present "
        "today, or what was paid last month. SedPayroll consolidates employee "
        "records, daily attendance, salary structures and monthly payroll processing "
        "into one internal web application, shared by every employee — administrators "
        "simply have a role flag unlocking extra sections, rather than a separate "
        "admin application or account type."
    )
    doc.add_heading("1.3 Scope of this document", level=2)
    doc.add_paragraph(
        "This BRD covers the business objectives, stakeholders, in-scope and "
        "out-of-scope business capabilities, numbered business requirements, "
        "assumptions, constraints, success criteria and known risks for SedPayroll "
        "as it stands today. It intentionally stays at the business-capability "
        "level; user-facing feature detail lives in the companion Product "
        "Requirements Document (PRD), and full functional/technical detail in the "
        "Software Requirements Specification (SRS)."
    )

    # 2. Objectives
    doc.add_heading("2. Business Objectives", level=1)
    add_bullets(doc, [
        "Maintain a single, authoritative system of record for employee, "
        "attendance and compensation data, replacing spreadsheets.",
        "Automate monthly payroll computation — including pro-rating pay for "
        "attendance — to remove manual calculation and its errors.",
        "Give every employee self-service visibility into their own attendance "
        "and payslips, reducing HR's day-to-day query load.",
        "Enforce role-based access to sensitive data (compensation, bank "
        "details, org-wide attendance) without operating a separate admin "
        "system or duplicate accounts.",
        "Give administrators an at-a-glance view of organizational headcount, "
        "attendance and payroll status.",
    ])

    # 3. Stakeholders
    doc.add_heading("3. Stakeholders", level=1)
    req_table(doc, ["Stakeholder", "Interest / role"], [
        ("HR / Payroll Administrator", "Primary system operator: manages employee "
         "records, runs monthly payroll, oversees org-wide attendance."),
        ("Employee (all staff)", "Self-service consumer: clocks in/out, views their "
         "own attendance history, payslips and profile."),
        ("Company leadership / Finance", "Consumes payroll output (net pay, run "
         "status) as the basis for actual salary disbursal, handled outside the "
         "system."),
        ("Engineering / IT", "Builds, deploys and maintains the application and "
         "its database."),
    ], col_widths=[5, 11])

    # 4. Business Scope
    doc.add_heading("4. Business Scope", level=1)
    doc.add_heading("4.1 In scope", level=2)
    req_table(doc, ["Business area", "Capability"], [
        ("Employee lifecycle", "Create, edit, deactivate/reactivate employees; "
         "grant or revoke administrative access."),
        ("Attendance", "Daily self-service clock-in/clock-out; org-wide daily "
         "attendance visibility for administrators."),
        ("Payroll processing", "Monthly payroll run per organization, computed "
         "from salary structures and attendance; mark runs paid."),
        ("Employee self-service", "View own attendance history, payslips and "
         "profile without contacting HR."),
        ("Access control", "Two-tier access (employee / administrator) on a "
         "single shared login, without separate admin infrastructure."),
    ], col_widths=[5, 11])

    doc.add_heading("4.2 Out of scope", level=2)
    add_bullets(doc, [
        "Multi-company / multi-tenant operation — the system serves one "
        "organization.",
        "Statutory compliance filing or remittance (EPFO/ESIC/TDS challans, "
        "returns) — the system records deduction figures only.",
        "Leave request / approval workflow — attendance carries an “on_leave” "
        "status, but there is no request-and-approve process to set it.",
        "Recruitment, onboarding checklists, or performance management.",
        "Actual salary disbursal / bank payment-gateway integration.",
        "Native mobile applications.",
        "Biometric or geo-fenced attendance capture.",
        "Email / SMS notifications of any kind.",
        "Multi-currency or multi-country payroll rules.",
    ])

    # 5. Business Requirements
    doc.add_heading("5. Business Requirements", level=1)

    doc.add_heading("5.1 Employee lifecycle management", level=2)
    req_table(doc, ["ID", "Requirement"], [
        ("BR-01", "The business shall maintain a single record per employee "
         "capturing personal, employment, bank and compensation details."),
        ("BR-02", "The system shall automatically provision a login for every "
         "employee record created, removing any separate account-creation step."),
        ("BR-03", "The business shall be able to grant or revoke administrative "
         "privileges on any employee's own login, rather than creating a "
         "separate administrator identity."),
        ("BR-04", "The business shall be able to mark an employee inactive on "
         "exit without losing their historical attendance or payroll records."),
    ], col_widths=[2.2, 13.8])

    doc.add_heading("5.2 Attendance", level=2)
    req_table(doc, ["ID", "Requirement"], [
        ("BR-05", "Every employee shall be able to record their own daily "
         "attendance (clock in / clock out) through self-service."),
        ("BR-06", "The business shall be able to view a consolidated daily "
         "attendance log across the organization for any given date."),
        ("BR-07", "Attendance shall automatically classify a short working day "
         "(under 5 worked hours) as a half day for payroll purposes."),
    ], col_widths=[2.2, 13.8])

    doc.add_heading("5.3 Payroll", level=2)
    req_table(doc, ["ID", "Requirement"], [
        ("BR-08", "The business shall be able to generate payroll for the whole "
         "organization for a given month in a single action, computed from each "
         "employee's current salary structure and recorded attendance."),
        ("BR-09", "Payroll shall automatically pro-rate pay for days not worked "
         "(loss of pay), based on recorded attendance, without manual "
         "calculation."),
        ("BR-10", "The business shall not be able to process payroll twice for "
         "the same month, to prevent duplicate disbursal figures."),
        ("BR-11", "Once payroll is processed, the business shall be able to mark "
         "it as paid to reflect actual disbursal status."),
    ], col_widths=[2.2, 13.8])

    doc.add_heading("5.4 Self-service & transparency", level=2)
    req_table(doc, ["ID", "Requirement"], [
        ("BR-12", "Every employee shall be able to view their own attendance "
         "history and payslips at any time, without contacting HR."),
        ("BR-13", "An employee's payslip and personal data shall never be "
         "visible to another employee."),
    ], col_widths=[2.2, 13.8])

    doc.add_heading("5.5 Access control", level=2)
    req_table(doc, ["ID", "Requirement"], [
        ("BR-14", "Only administrators shall be able to manage employees, run "
         "payroll, and view organization-wide attendance."),
        ("BR-15", "The system shall prevent the organization from ever being "
         "left with zero administrators."),
    ], col_widths=[2.2, 13.8])

    # 6. Assumptions & Constraints
    doc.add_heading("6. Assumptions & Constraints", level=1)
    add_bullets(doc, [
        "The organization is a single legal entity; no multi-tenant separation "
        "is required.",
        "The business hosts its own deployment and PostgreSQL database "
        "(self-hosted Docker Compose locally, or a managed Postgres such as "
        "Vercel Postgres in production).",
        "Statutory payroll compliance (PF/ESI remittance, TDS computation and "
        "filing) is handled by the business outside this system; SedPayroll "
        "tracks the deduction amounts only, not the statutory rules behind "
        "them.",
        "New employees are given a fixed default temporary password "
        "(communicated by HR outside the system); there is no automated email "
        "delivery.",
        "All figures are in Indian Rupees; no multi-currency support is "
        "assumed.",
    ])

    # 7. Success criteria
    doc.add_heading("7. Success Criteria / KPIs", level=1)
    add_bullets(doc, [
        "Monthly payroll can be run in a single action in under 5 minutes, "
        "down from a manual spreadsheet process.",
        "Zero duplicate payroll runs per calendar month (enforced by the "
        "system, not by process discipline).",
        "100% of employees can retrieve their own payslip without raising an "
        "HR request.",
        "Measurable reduction in HR queries related to attendance/payslip "
        "status after rollout.",
    ])

    # 8. Risks
    doc.add_heading("8. Risks & Mitigations", level=1)
    req_table(doc, ["Risk", "Impact", "Mitigation"], [
        ("Default temporary passwords (Employee@123 / Admin@123) are fixed and "
         "predictable.", "Credential-guessing risk on any freshly created "
         "account.", "Recommend forcing a password change on first login "
         "(not yet implemented — tracked as a future requirement)."),
        ("No leave-request/approval workflow exists.", "The “on_leave” "
         "attendance status has no way to be set through the current UI.",
         "Scope a leave-management module, or an admin action to set "
         "attendance status directly."),
        ("No audit trail of who changed what.", "Limited accountability for "
         "sensitive changes (salary, admin-role grants).", "Add an audit log "
         "as a future enhancement."),
        ("No statutory compliance calculation.", "PF/PT figures entered are not "
         "validated against government slabs.", "Business must verify "
         "statutory correctness externally; out of scope by design."),
    ], col_widths=[5, 5, 6])

    # 9. Glossary
    doc.add_heading("9. Glossary", level=1)
    req_table(doc, ["Term", "Meaning"], [
        ("CTC", "Cost to Company — the fully-loaded annual compensation figure "
         "stored alongside a salary structure."),
        ("HRA", "House Rent Allowance — a standard Indian salary component."),
        ("PF", "Provident Fund — a statutory retirement-savings deduction."),
        ("PT", "Professional Tax — a state-level statutory deduction."),
        ("Loss of pay (LOP)", "The portion of gross pay withheld for days not "
         "worked in a given month."),
        ("Payroll run", "One month's payroll processing event for the whole "
         "organization; produces one payslip per active employee."),
    ], col_widths=[4, 12])

    add_page_number_footer(doc)
    return doc


# ---------------------------------------------------------------------------
# PRD
# ---------------------------------------------------------------------------

def build_prd():
    doc = new_document()
    add_cover_page(
        doc,
        doc_type="PRD",
        title="Product Requirements Document",
        subtitle="Payroll, attendance and employee management — features, personas and scope",
        version=VERSION,
        date_str=DATE_STR,
        status=STATUS,
    )
    add_document_control(doc, history_rows=HISTORY_ROWS)
    add_toc(doc)
    doc.add_page_break()

    doc.add_heading("1. Product Overview", level=1)
    doc.add_paragraph(
        "SedPayroll is one shared web portal for payroll, attendance and "
        "employee management. Every person who signs in reaches the same "
        "application — administrator is simply a role flag on an employee's "
        "own login, unlocking three extra sections (Employees, Team "
        "Attendance, Payroll), rather than a separate admin product or "
        "account."
    )
    p = doc.add_paragraph()
    p.add_run("Product pillars: ").bold = True
    p.add_run(
        "Self-service (employees never need to ask HR for their own data), "
        "Simplicity (one portal, one login, no parallel admin app), and "
        "Trust (role-scoped access enforced consistently, not just in the UI)."
    )

    doc.add_heading("2. Target Users & Personas", level=1)
    doc.add_heading("2.1 Priya — HR / Payroll Administrator", level=3)
    add_bullets(doc, [
        "Goal: onboard new hires, keep the roster accurate, and close monthly "
        "payroll quickly and correctly.",
        "Pain today: manual spreadsheet payroll calculation is slow and "
        "error-prone; answering “what's my attendance/payslip” questions "
        "eats into her day.",
        "How SedPayroll helps: one action runs payroll for the whole "
        "organization; employees self-serve their own attendance and "
        "payslips.",
    ])
    doc.add_heading("2.2 Rohan — Employee", level=3)
    add_bullets(doc, [
        "Goal: clock in/out reliably, and check his attendance and pay "
        "without emailing HR.",
        "Pain today: no visibility into his own attendance record or when "
        "payslips are ready.",
        "How SedPayroll helps: a dashboard clock widget, a full attendance "
        "history, and payslips available the moment payroll is processed.",
    ])

    doc.add_heading("3. Goals & Non-Goals", level=1)
    doc.add_heading("3.1 Goals", level=2)
    add_bullets(doc, [
        "One login, one portal, for every employee regardless of role.",
        "Attendance-aware payroll: pay is pro-rated automatically from "
        "recorded attendance, with no manual adjustment step.",
        "Zero-configuration self-service: every employee can see their own "
        "data the moment it exists, no setup required.",
        "Guardrails baked in, not bolted on: duplicate payroll runs and "
        "zero-admin lockouts are prevented by the system itself.",
    ])
    doc.add_heading("3.2 Non-goals (explicitly out of scope for this product)", level=2)
    add_bullets(doc, [
        "Leave request/approval workflows.",
        "Multi-tenant / multi-company support.",
        "Statutory compliance filing (EPFO, ESIC, TDS returns).",
        "Native mobile applications.",
        "Email / SMS / push notifications.",
        "Biometric or location-based attendance capture.",
        "Bank disbursal / payment-gateway integration.",
        "Multi-currency or multi-country payroll.",
    ])

    doc.add_heading("4. Information Architecture", level=1)
    doc.add_paragraph(
        "Navigation is a single list that grows by role rather than a "
        "separate admin shell:"
    )
    req_table(doc, ["Section", "Route", "Visible to"], [
        ("Dashboard", "/", "Everyone"),
        ("My Attendance", "/attendance", "Everyone"),
        ("My Payslips", "/payslips", "Everyone"),
        ("Profile", "/profile", "Everyone"),
        ("Employees", "/employees", "Administrators only"),
        ("Team Attendance", "/team-attendance", "Administrators only"),
        ("Payroll", "/payroll", "Administrators only"),
    ], col_widths=[4, 5, 7])

    doc.add_heading("5. Features & User Stories", level=1)
    doc.add_paragraph(
        "Every story below is already implemented (v1.0); priority reflects "
        "how essential the capability is to the product, using MoSCoW."
    )

    def feature_module(title, stories):
        doc.add_heading(title, level=2)
        req_table(doc, ["User story", "Acceptance criteria", "Priority"], stories,
                  col_widths=[6, 8, 2])

    feature_module("5.1 Authentication", [
        ("As any user, I want to sign in with my work email and password so "
         "that I can reach my portal.",
         "Valid credentials redirect to the dashboard; invalid credentials "
         "show one generic error without revealing which field was wrong.",
         "Must"),
        ("As a signed-in user, I want an admin-only page I try to open to "
         "quietly send me back to my dashboard, not show an error.",
         "Non-admin requests to /employees, /team-attendance or /payroll "
         "(and their sub-paths) redirect to “/”.",
         "Must"),
        ("As a user, I want to sign out.",
         "Signing out clears the session and returns me to /login.",
         "Must"),
    ])

    feature_module("5.2 Dashboard", [
        ("As an employee, I want to clock in/out from my dashboard so I don't "
         "need a separate attendance screen.",
         "A widget shows today's status and disables the button that no "
         "longer applies (e.g. Clock in once already clocked in).",
         "Must"),
        ("As an employee, I want to see my latest payslip's net pay at a "
         "glance.",
         "Shows the most recent payslip's net pay, or an empty-state message "
         "before any payroll has run for me.",
         "Must"),
        ("As an admin, I want an organization snapshot on my dashboard.",
         "Shows active/inactive employee counts, employees present today, and "
         "payroll run count, each linking to its section; lists the 5 most "
         "recent payroll runs.",
         "Must"),
    ])

    feature_module("5.3 My Attendance", [
        ("As an employee, I want to clock in once per day.",
         "A second clock-in the same day is blocked (button disabled; the "
         "underlying action also rejects it).",
         "Must"),
        ("As an employee, I want to clock out and have a fair status "
         "recorded.",
         "Worked hours under 5 record a half day; 5+ records a full present "
         "day.",
         "Must"),
        ("As an employee, I want to review my attendance history.",
         "Shows up to my last 60 records, most recent first.",
         "Must"),
    ])

    feature_module("5.4 Team Attendance (admin)", [
        ("As an admin, I want to see everyone's attendance for a chosen day.",
         "Defaults to today; every active employee is listed even with no "
         "record (shown as unmarked).",
         "Must"),
        ("As an admin, I want to jump to any date and share that view.",
         "The selected date is reflected in the URL as a query parameter.",
         "Should"),
    ])

    feature_module("5.5 Employees (admin)", [
        ("As an admin, I want to see active employees by default and switch "
         "to inactive.",
         "Active/Inactive tabs filter the roster; empty state shown when a "
         "tab has no rows.",
         "Must"),
        ("As an admin, I want adding an employee to also set them up to log "
         "in, with no extra step.",
         "Creating an employee auto-generates an employee code, a login with "
         "a default temporary password, and an initial salary structure, all "
         "in one action.",
         "Must"),
        ("As an admin, I want to be stopped from creating a duplicate "
         "account.",
         "An email already used by another login is rejected before any row "
         "is written.",
         "Must"),
        ("As an admin, I want to edit an employee's details and update their "
         "compensation without losing history.",
         "Edits validate the same required fields; a new salary-structure "
         "record is added rather than overwriting the previous one.",
         "Must"),
        ("As an admin, I want to offboard someone without losing their "
         "history, and rehire them later.",
         "Deactivate sets status=inactive and records a leaving date; "
         "reactivate reverses it. Historical attendance/payroll data is "
         "untouched either way.",
         "Must"),
        ("As an admin, I want to grant or remove another admin's access "
         "safely.",
         "I can't remove my own admin access, and the system won't let the "
         "last remaining admin be demoted.",
         "Must"),
    ])

    feature_module("5.6 Payroll (admin)", [
        ("As an admin, I want to run payroll for the whole organization for a "
         "month in one action.",
         "One “Run payroll” action computes and stores a payslip for every "
         "active employee for the chosen month/year.",
         "Must"),
        ("As an admin, I want the system to stop me running payroll twice for "
         "the same month.",
         "A repeat run for an already-processed month/year is rejected with a "
         "clear message.",
         "Must"),
        ("As an admin, I want pay automatically adjusted for attendance.",
         "Gross pay is pro-rated by days present ÷ days in month (half days "
         "count as 0.5, paid leave counts as a full day); the shortfall shows "
         "as loss-of-pay.",
         "Must"),
        ("As an admin, I want to review a run before treating it as "
         "disbursed.",
         "A run's detail page lists every payslip with days present, gross, "
         "deductions and net pay, plus the run's total.",
         "Must"),
        ("As an admin, I want to record that a run has actually been paid "
         "out.",
         "“Mark as paid” is available only while a run is “processed”, and "
         "moves it to “paid”.",
         "Must"),
    ])

    feature_module("5.7 My Payslips", [
        ("As an employee, I want to see all my payslips.",
         "Lists every payslip I own, most recent month first, with gross and "
         "net pay.",
         "Must"),
        ("As an employee, I want to see exactly how my pay was calculated.",
         "Opening a payslip shows the full earnings and deductions breakdown "
         "and the final net pay.",
         "Must"),
        ("As an employee, I must never be able to see someone else's "
         "payslip, even by guessing a URL.",
         "Any payslip id that isn't mine returns “not found”.",
         "Must"),
    ])

    feature_module("5.8 Profile", [
        ("As an employee, I want to see my own details on file.",
         "Read-only view of code, name, email, phone, department, "
         "designation, join date, birth date and address.",
         "Must"),
        ("As an employee, if I need something changed I want to know who to "
         "ask.",
         "The page tells me to contact HR — there is no self-edit form.",
         "Won't (by design, for now)"),
    ])

    doc.add_heading("6. Design Principles", level=1)
    add_bullets(doc, [
        "One shared portal: admin sections are additive navigation items, not "
        "a separate application shell or login.",
        "Light and dark themes, following the system preference or an "
        "explicit user toggle.",
        "An indigo/emerald brand palette used consistently for primary "
        "actions and positive states (active, paid, present).",
        "Every interactive element carries a stable identifier to keep the "
        "product testable as it grows (see the companion automated test "
        "suite).",
    ])

    doc.add_heading("7. Success Metrics", level=1)
    add_bullets(doc, [
        "% of employees clocking in/out through the portal daily.",
        "Time taken to complete a monthly payroll run, start to “marked "
        "paid”.",
        "Reduction in HR-directed queries about attendance/payslip status.",
        "Zero incidents of duplicate payroll runs or zero-admin lockouts.",
    ])

    doc.add_heading("8. Release Plan / Phasing", level=1)
    doc.add_heading("8.1 v1.0 — current", level=2)
    doc.add_paragraph("Every module in Section 5, as shipped today.")
    doc.add_heading("8.2 Candidate v1.1 (proposed, not committed)", level=2)
    add_bullets(doc, [
        "Force a password change on first login.",
        "Configurable half-day threshold (currently a fixed 5 hours).",
        "A leave-request/approval workflow feeding attendance status.",
        "An audit log of administrative changes.",
        "Limited employee self-edit (e.g. phone, address).",
        "Email notifications (payslip ready, admin access changed, etc.).",
    ])

    doc.add_heading("9. Open Questions", level=1)
    add_numbered(doc, [
        "Should employees be able to update limited contact details "
        "themselves?",
        "Should leave requests be a first-class workflow distinct from raw "
        "attendance status?",
        "Should payroll require a second approver (maker-checker) before "
        "being marked paid?",
        "Should the half-day attendance threshold be configurable per "
        "organization?",
    ])

    doc.add_heading("10. Appendix — Test Coverage", level=1)
    doc.add_paragraph(
        "Every story above is covered by the automated Playwright suite under "
        "tests/e2e/, with the full manual/automated test-case inventory in "
        "docs/SedPayroll-E2E-Test-Cases.xlsx."
    )

    add_page_number_footer(doc)
    return doc


# ---------------------------------------------------------------------------
# SRS
# ---------------------------------------------------------------------------

def build_srs():
    doc = new_document()
    add_cover_page(
        doc,
        doc_type="SRS",
        title="Software Requirements Specification",
        subtitle="Functional, data and non-functional requirements, IEEE 830 style",
        version=VERSION,
        date_str=DATE_STR,
        status=STATUS,
    )
    add_document_control(doc, history_rows=HISTORY_ROWS)
    add_toc(doc)
    doc.add_page_break()

    # 1. Introduction
    doc.add_heading("1. Introduction", level=1)
    doc.add_heading("1.1 Purpose", level=2)
    doc.add_paragraph(
        "This Software Requirements Specification (SRS) describes the "
        "functional, data and non-functional requirements of SedPayroll in "
        "implementation-level detail, for engineers, QA and technical "
        "reviewers. It documents the system as currently built."
    )
    doc.add_heading("1.2 Scope", level=2)
    doc.add_paragraph(
        "SedPayroll is a single-tenant web application covering employee "
        "records, daily attendance, salary structures and monthly payroll "
        "processing, shared by every employee of one organization with "
        "role-based access (employee / administrator)."
    )
    doc.add_heading("1.3 Definitions, Acronyms and Abbreviations", level=2)
    req_table(doc, ["Term", "Definition"], [
        ("RBAC", "Role-Based Access Control."),
        ("JWT", "JSON Web Token — the signed, stateless session format used "
         "by Auth.js."),
        ("HRA / CTC / PF / PT", "House Rent Allowance / Cost to Company / "
         "Provident Fund / Professional Tax — standard Indian payroll "
         "terms."),
        ("LOP", "Loss of Pay — pay withheld for days not worked in a month."),
        ("FR", "Functional Requirement (this document's numbering)."),
        ("Server Action", "A Next.js App Router mechanism for handling form "
         "submissions/mutations directly from a Server Component, without a "
         "separate REST endpoint."),
    ], col_widths=[4, 12])
    doc.add_heading("1.4 References", level=2)
    add_bullets(doc, [
        "SedPayroll source repository (README.md, AGENTS.md, db/schema.ts, "
        "lib/, app/).",
        "SedPayroll BRD v1.0 and PRD v1.0 (companion documents).",
        "docs/SedPayroll-E2E-Test-Cases.xlsx — the test-case inventory.",
    ])
    doc.add_heading("1.5 Overview", level=2)
    doc.add_paragraph(
        "Section 2 describes the system at a high level; Section 3 lists "
        "detailed functional requirements by module; Section 4 covers "
        "external interfaces; Section 5 the data model; Section 6 "
        "non-functional requirements; Section 7 other requirements; "
        "appendices follow."
    )

    # 2. Overall description
    doc.add_heading("2. Overall Description", level=1)
    doc.add_heading("2.1 Product perspective", level=2)
    doc.add_paragraph(
        "SedPayroll is a standalone Next.js 16 (App Router) application "
        "backed by a single PostgreSQL database, using Server Actions for "
        "all mutations instead of a conventional REST/JSON API. "
        "Authentication is provided by Auth.js v5 (Credentials provider, "
        "JWT session strategy)."
    )
    doc.add_heading("2.2 Product functions (summary)", level=2)
    add_bullets(doc, [
        "Authenticate users and enforce role-based route access.",
        "Record and display daily attendance (self-service and org-wide).",
        "Maintain employee master data and versioned salary structures.",
        "Run monthly payroll, pro-rating pay from attendance, and track its "
        "status through to “paid”.",
        "Present payslips to their owning employee only.",
    ])
    doc.add_heading("2.3 User classes and characteristics", level=2)
    req_table(doc, ["Class", "Characteristics"], [
        ("Employee", "Any authenticated user with role=employee. Full "
         "self-service access to their own dashboard, attendance, payslips "
         "and profile."),
        ("Administrator", "Any authenticated user with role=admin — the same "
         "kind of account as Employee, with an additional role flag. Gains "
         "the Employees, Team Attendance and Payroll sections on top of the "
         "employee capabilities above."),
    ], col_widths=[3, 13])
    doc.add_heading("2.4 Operating environment", level=2)
    add_bullets(doc, [
        "Server runtime: Node.js running Next.js 16 (Turbopack), deployable "
        "to Vercel or any Node-compatible host.",
        "Database: PostgreSQL 16 — local via Docker Compose in development, "
        "or a managed Postgres (e.g. Vercel Postgres) in production.",
        "Client: any evergreen desktop or mobile browser; the automated test "
        "suite exercises Chromium, and the toolchain also supports Firefox "
        "and WebKit.",
    ])
    doc.add_heading("2.5 Design and implementation constraints", level=2)
    add_bullets(doc, [
        "All mutations are implemented as Next.js Server Actions bound to "
        "forms — there is no separate versioned REST/JSON API beyond the "
        "Auth.js callback routes.",
        "Drizzle ORM is the sole data-access layer; the schema in "
        "db/schema.ts is the single source of truth for the relational "
        "model.",
        "The system is single-tenant: one organization's data per "
        "deployment/database.",
    ])
    doc.add_heading("2.6 Assumptions and dependencies", level=2)
    add_bullets(doc, [
        "DB_DATABASE_URL and AUTH_SECRET environment variables are supplied "
        "for every environment.",
        "bcryptjs is used for password hashing; postgres-js is the Postgres "
        "driver.",
        "Statutory payroll compliance is out of scope (see BRD §4.2); PF/PT "
        "are user-entered deduction amounts, not computed against government "
        "slabs.",
    ])

    # 3. Functional requirements
    doc.add_heading("3. System Features (Functional Requirements)", level=1)

    doc.add_heading("3.1 Authentication & session (FR-AUTH)", level=2)
    req_table(doc, ["ID", "Requirement"], [
        ("FR-AUTH-01", "The system shall authenticate a user by comparing the "
         "submitted password against the bcrypt hash stored for the matching "
         "email in the users table."),
        ("FR-AUTH-02", "On success the system shall issue a JWT session "
         "embedding the user's role and employeeId."),
        ("FR-AUTH-03", "Login failure — unknown email or wrong password — "
         "shall present a single generic message (“Invalid email or "
         "password.”) without indicating which factor failed."),
        ("FR-AUTH-04", "An authenticated user requesting /login shall be "
         "redirected to the dashboard (“/”)."),
        ("FR-AUTH-05", "An unauthenticated request to any portal route shall "
         "be redirected to /login."),
        ("FR-AUTH-06", "Signing out shall terminate the session and redirect "
         "to /login."),
        ("FR-AUTH-07", "Requests to /employees, /team-attendance or /payroll "
         "(and any sub-paths) from a non-admin session shall be redirected "
         "to “/” rather than shown an error."),
        ("FR-AUTH-08", "Route protection shall be enforced at two "
         "independent layers: edge middleware for page navigations, and a "
         "per-request guard inside every server action/page, since Server "
         "Actions are directly reachable and bypass the edge matcher."),
    ], col_widths=[2.4, 13.6])

    doc.add_heading("3.2 Dashboard (FR-DASH)", level=2)
    req_table(doc, ["ID", "Requirement"], [
        ("FR-DASH-01", "The dashboard shall display the signed-in user's "
         "clock-in/out widget reflecting today's attendance state."),
        ("FR-DASH-02", "The dashboard shall display the employee's most "
         "recently generated payslip's net pay, or an empty-state message if "
         "none exists."),
        ("FR-DASH-03", "For an administrator only, the dashboard shall "
         "additionally show: active employee count, inactive employee "
         "count, count of employees present or half-day today, and total "
         "payroll run count — each linking to its section."),
        ("FR-DASH-04", "The administrator's dashboard shall list the 5 most "
         "recently created payroll runs with month, status and processed "
         "date."),
    ], col_widths=[2.4, 13.6])

    doc.add_heading("3.3 Attendance — self-service (FR-ATT)", level=2)
    req_table(doc, ["ID", "Requirement"], [
        ("FR-ATT-01", "An employee may clock in at most once per calendar "
         "day; a second attempt the same day shall be rejected with “You "
         "have already clocked in today.” and no duplicate row created."),
        ("FR-ATT-02", "Clocking in shall record the current server timestamp "
         "as clockIn and set status=present."),
        ("FR-ATT-03", "An employee may clock out only after clocking in the "
         "same day; clocking out with no prior clock-in shall be rejected "
         "with “You need to clock in first.”"),
        ("FR-ATT-04", "A second clock-out attempt the same day shall be "
         "rejected with “You have already clocked out today.”"),
        ("FR-ATT-05", "On clock-out, the system shall compute elapsed hours "
         "(clockOut − clockIn); if the result is under 5 hours, status shall "
         "be set to half_day, otherwise present."),
        ("FR-ATT-06", "An employee shall be able to view up to their most "
         "recent 60 attendance records, ordered most recent first."),
    ], col_widths=[2.4, 13.6])

    doc.add_heading("3.4 Team attendance — administrator (FR-TATT)", level=2)
    req_table(doc, ["ID", "Requirement"], [
        ("FR-TATT-01", "An administrator shall be able to view a single "
         "day's attendance for every active employee, defaulting to the "
         "current date."),
        ("FR-TATT-02", "The administrator shall be able to select any date "
         "via a date picker; the selected date shall be reflected in the "
         "page URL as a query parameter."),
        ("FR-TATT-03", "An active employee with no attendance row for the "
         "selected date shall be displayed with unmarked (“—”) clock-in / "
         "clock-out cells, not treated as absent by default."),
    ], col_widths=[2.4, 13.6])

    doc.add_heading("3.5 Employee management — administrator (FR-EMP)", level=2)
    req_table(doc, ["ID", "Requirement"], [
        ("FR-EMP-01", "The employee list shall be filterable by status "
         "(active, default; or inactive), showing employee code, name, "
         "email, department, designation, date of joining, status and "
         "role."),
        ("FR-EMP-02", "Creating an employee shall require first name, last "
         "name, a unique email address, date of joining and a basic salary "
         "amount (≥ 0); department, designation, gender, date of birth, "
         "address and bank details are optional."),
        ("FR-EMP-03", "On successful creation, the system shall, within a "
         "single database transaction: assign a sequential employee code "
         "(“EMP” + zero-padded running count), create a login with a "
         "default password (“Admin@123” if admin access is granted at "
         "creation, otherwise “Employee@123”), and create an initial "
         "salary-structure record effective from the date of joining."),
        ("FR-EMP-04", "Creating or editing an employee with an email already "
         "used by a different employee's login shall be rejected with a "
         "field-level error before any row is written or changed."),
        ("FR-EMP-05", "Editing an employee shall re-validate the same "
         "required fields as creation and insert a new salary-structure row "
         "effective from the current date, preserving all prior structures "
         "for historical payroll accuracy, rather than overwriting the "
         "previous record."),
        ("FR-EMP-06", "An administrator shall be able to deactivate an "
         "active employee (sets status=inactive and dateOfLeaving=today) and "
         "reactivate an inactive one (sets status=active and clears "
         "dateOfLeaving); inactive employees shall be excluded from active "
         "counts, team-attendance listings and future payroll runs, while "
         "retaining all historical attendance and payslip records."),
        ("FR-EMP-07", "An administrator shall be able to grant or revoke "
         "admin access on any employee's login, subject to: (a) an "
         "administrator may not remove their own admin access; (b) the "
         "system shall refuse to demote an administrator while the total "
         "admin count is ≤ 1."),
        ("FR-EMP-08", "Immediately after a successful creation, the system "
         "shall display the new employee's temporary password once, for the "
         "administrator to relay to them."),
    ], col_widths=[2.4, 13.6])

    doc.add_heading("3.6 Payroll — administrator (FR-PAY)", level=2)
    req_table(doc, ["ID", "Requirement"], [
        ("FR-PAY-01", "An administrator shall be able to run payroll for a "
         "specified month (1–12) and calendar year."),
        ("FR-PAY-02", "The system shall reject running payroll for a "
         "(month, year) pair that already has a run, with the message "
         "“Payroll for this month has already been run.”"),
        ("FR-PAY-03", "Running payroll shall generate exactly one payslip "
         "for every currently active employee, using that employee's latest "
         "salary structure with effectiveFrom on or before the last day of "
         "the target month."),
        ("FR-PAY-04", "An employee's pro-rated gross for the month shall "
         "equal (basic + HRA + conveyance + medical allowance + special "
         "allowance) × min(daysPresent ÷ daysInMonth, 1), where "
         "daysPresent sums 1 per “present” day, 0.5 per “half_day”, 1 per "
         "“on_leave” day (treated as fully paid), and 0 for any day with no "
         "attendance record or status “absent”."),
        ("FR-PAY-05", "Loss-of-pay shall equal the full (non-prorated) gross "
         "minus the pro-rated gross. Total deductions shall equal provident "
         "fund + professional tax + loss-of-pay. Net pay shall equal the "
         "pro-rated gross minus provident fund minus professional tax."),
        ("FR-PAY-06", "Each generated payslip shall store a structured "
         "breakdown of its individually pro-rated earnings components "
         "(basic, HRA, conveyance, medical allowance, special allowance) and "
         "its deductions (provident fund, professional tax, loss-of-pay)."),
        ("FR-PAY-07", "A payroll run shall carry a status of draft, "
         "processed or paid; a run is created directly in “processed” "
         "status once computed, and “Mark as paid” shall only be available "
         "while status is “processed”, transitioning it to “paid”."),
        ("FR-PAY-08", "An administrator shall be able to open any payroll "
         "run to view every payslip in it (days present, gross earnings, "
         "deductions, net pay) alongside the run's total net pay."),
        ("FR-PAY-09", "Requesting a payroll run detail page for a "
         "non-existent run id shall return a not-found (HTTP 404) response."),
    ], col_widths=[2.4, 13.6])

    doc.add_heading("3.7 Payslips — self-service (FR-PSLIP)", level=2)
    req_table(doc, ["ID", "Requirement"], [
        ("FR-PSLIP-01", "An employee shall be able to list every payslip "
         "belonging to them, ordered most recent payroll month first, "
         "showing month, gross earnings and net pay."),
        ("FR-PSLIP-02", "An employee shall be able to open one of their "
         "payslips to view the full earnings breakdown, deductions "
         "breakdown, and net pay."),
        ("FR-PSLIP-03", "A request for a payslip id that does not belong to "
         "the requesting employee's own employeeId shall return a "
         "not-found (HTTP 404) response, regardless of the requester's role, "
         "so as not to confirm the id's existence."),
    ], col_widths=[2.4, 13.6])

    doc.add_heading("3.8 Profile (FR-PROF)", level=2)
    req_table(doc, ["ID", "Requirement"], [
        ("FR-PROF-01", "An employee shall be able to view a read-only "
         "summary of their own record: employee code, full name, email, "
         "phone, department, designation, date of joining, date of birth "
         "and address."),
        ("FR-PROF-02", "The profile page shall present no editable form "
         "controls for any field."),
    ], col_widths=[2.4, 13.6])

    # 4. External interfaces
    doc.add_heading("4. External Interface Requirements", level=1)
    doc.add_heading("4.1 User interfaces", level=2)
    doc.add_paragraph(
        "A responsive web UI: a fixed sidebar navigation on desktop widths, "
        "collapsing to a horizontal top navigation on narrow/mobile widths. "
        "Light and dark themes are supported, following the OS preference by "
        "default with an explicit user toggle, built with Tailwind CSS and "
        "shadcn/ui component primitives."
    )
    doc.add_heading("4.2 Hardware interfaces", level=2)
    doc.add_paragraph("None beyond a standard web client (desktop or mobile).")
    doc.add_heading("4.3 Software interfaces", level=2)
    add_bullets(doc, [
        "PostgreSQL 16, accessed via the postgres-js driver through Drizzle "
        "ORM.",
        "Auth.js v5 (Credentials provider, JWT session strategy) for "
        "authentication.",
        "No third-party payment, email, or SMS integrations.",
    ])
    doc.add_heading("4.4 Communications interfaces", level=2)
    doc.add_paragraph(
        "HTTPS in production; plain HTTP is used only for local development."
    )

    # 5. Data requirements
    doc.add_heading("5. Data Requirements", level=1)
    doc.add_paragraph(
        "The relational schema (db/schema.ts, Drizzle ORM / PostgreSQL) "
        "comprises eight tables:"
    )

    doc.add_heading("5.1 departments", level=3)
    req_table(doc, ["Field", "Type / constraint"], [
        ("id", "serial, primary key"),
        ("name", "varchar(120), unique, required"),
    ], col_widths=[4, 12])

    doc.add_heading("5.2 designations", level=3)
    req_table(doc, ["Field", "Type / constraint"], [
        ("id", "serial, primary key"),
        ("title", "varchar(120), unique, required"),
    ], col_widths=[4, 12])

    doc.add_heading("5.3 employees", level=3)
    req_table(doc, ["Field", "Type / constraint"], [
        ("id", "serial, primary key"),
        ("employeeCode", "varchar(20), unique, required"),
        ("firstName / lastName", "varchar(100), required"),
        ("email", "varchar(255), unique, required"),
        ("phone", "varchar(20), optional"),
        ("departmentId / designationId", "FK → departments/designations, "
         "ON DELETE SET NULL"),
        ("gender", "enum: male | female | other, optional"),
        ("dateOfBirth", "date, optional"),
        ("dateOfJoining", "date, required"),
        ("dateOfLeaving", "date, set on deactivation"),
        ("status", "enum: active | inactive, default active"),
        ("address, bankAccountNumber, bankName, ifsc, panNumber",
         "text/varchar, optional"),
        ("createdAt / updatedAt", "timestamptz"),
    ], col_widths=[5, 11])

    doc.add_heading("5.4 users", level=3)
    req_table(doc, ["Field", "Type / constraint"], [
        ("id", "serial, primary key"),
        ("email", "varchar(255), unique, required"),
        ("passwordHash", "text, required (bcrypt)"),
        ("role", "enum: admin | employee, default employee"),
        ("employeeId", "FK → employees, unique (1:1), ON DELETE CASCADE"),
        ("createdAt", "timestamptz"),
    ], col_widths=[4, 12])

    doc.add_heading("5.5 attendance", level=3)
    req_table(doc, ["Field", "Type / constraint"], [
        ("id", "serial, primary key"),
        ("employeeId", "FK → employees, ON DELETE CASCADE"),
        ("date", "date, required"),
        ("clockIn / clockOut", "timestamptz, nullable"),
        ("status", "enum: present | absent | half_day | on_leave, default "
         "present"),
        ("notes", "text, optional"),
        ("(employeeId, date)", "unique index — one row per employee per "
         "day"),
    ], col_widths=[4, 12])

    doc.add_heading("5.6 salary_structures", level=3)
    req_table(doc, ["Field", "Type / constraint"], [
        ("id", "serial, primary key"),
        ("employeeId", "FK → employees, ON DELETE CASCADE"),
        ("effectiveFrom", "date, required — versions the structure"),
        ("basic", "numeric(12,2), required"),
        ("hra, conveyance, medicalAllowance, specialAllowance, "
         "providentFund, professionalTax", "numeric(12,2), default 0"),
        ("ctc", "numeric(12,2), required"),
        ("createdAt", "timestamptz"),
    ], col_widths=[4, 12])

    doc.add_heading("5.7 payroll_runs", level=3)
    req_table(doc, ["Field", "Type / constraint"], [
        ("id", "serial, primary key"),
        ("month", "integer 1–12, required"),
        ("year", "integer, required"),
        ("status", "enum: draft | processed | paid, default draft"),
        ("processedAt", "timestamptz, nullable"),
        ("processedBy", "FK → users, ON DELETE SET NULL"),
        ("(month, year)", "unique index — one run per calendar month"),
    ], col_widths=[4, 12])

    doc.add_heading("5.8 payslips", level=3)
    req_table(doc, ["Field", "Type / constraint"], [
        ("id", "serial, primary key"),
        ("payrollRunId", "FK → payroll_runs, ON DELETE CASCADE"),
        ("employeeId", "FK → employees, ON DELETE CASCADE"),
        ("daysPresent", "numeric(5,1), required"),
        ("daysInMonth", "integer, required"),
        ("grossEarnings, totalDeductions, netPay", "numeric(12,2), "
         "required"),
        ("breakdown", "jsonb — structured earnings/deductions detail "
         "(FR-PAY-06)"),
        ("generatedAt", "timestamptz"),
    ], col_widths=[4, 12])

    doc.add_heading("5.9 Referential integrity summary", level=3)
    doc.add_paragraph(
        "Deleting an employee cascades to its login, attendance, salary "
        "structures and payslips (all ON DELETE CASCADE). Deleting a "
        "department or designation sets the employee's corresponding field "
        "to null rather than deleting the employee (ON DELETE SET NULL). "
        "Deleting a payroll run cascades to its payslips. In normal product "
        "use, employees are deactivated rather than deleted — hard deletion "
        "of these rows is not exposed anywhere in the application UI."
    )

    # 6. Non-functional
    doc.add_heading("6. Non-Functional Requirements", level=1)
    doc.add_heading("6.1 Performance", level=2)
    doc.add_paragraph(
        "No hard SLA is currently defined. Interactive pages are expected to "
        "respond in well under a second against a roster of tens to low "
        "hundreds of employees. Running payroll processes active employees "
        "sequentially (one DB round-trip per employee); this is a known "
        "scalability limit for very large rosters and is not currently "
        "batched or parallelized."
    )
    doc.add_heading("6.2 Security", level=2)
    add_bullets(doc, [
        "Passwords are stored as bcrypt hashes (cost factor 10); plaintext "
        "passwords are never persisted.",
        "Sessions are signed, stateless JWTs (AUTH_SECRET); there is no "
        "server-side session store to revoke individual sessions early.",
        "Every admin-only server action independently re-verifies the "
        "caller's role, in addition to edge-middleware route protection "
        "(defense in depth).",
        "Known gap: default temporary passwords are fixed, predictable "
        "strings (“Employee@123” / “Admin@123”) and are not forced to change "
        "on first login — tracked in the BRD's risk register.",
    ])
    doc.add_heading("6.3 Usability", level=2)
    doc.add_paragraph(
        "Every interactive element carries a stable data-testid attribute to "
        "support automated end-to-end testing. The login screen surfaces "
        "demo credentials directly, which is appropriate for evaluation/demo "
        "environments but should be removed or gated before a production "
        "rollout."
    )
    doc.add_heading("6.4 Reliability / availability", level=2)
    doc.add_paragraph(
        "No formal uptime SLA is defined; this is an internal tool that "
        "relies on the availability of its hosting platform (Vercel and/or "
        "managed Postgres) rather than any bespoke high-availability "
        "architecture."
    )
    doc.add_heading("6.5 Maintainability", level=2)
    doc.add_paragraph(
        "The codebase is TypeScript end-to-end with Drizzle-inferred types "
        "from a single schema source of truth; schema changes are tracked as "
        "versioned migrations under drizzle/."
    )
    doc.add_heading("6.6 Portability", level=2)
    doc.add_paragraph(
        "The only environment-specific dependency is a standard PostgreSQL "
        "connection string, so the same codebase runs unmodified against a "
        "local Docker Compose Postgres or a managed Postgres provider."
    )
    doc.add_heading("6.7 Accessibility", level=2)
    doc.add_paragraph(
        "Semantic heading/table structure and a “skip to main content” link "
        "are present. No formal WCAG conformance audit has been performed; "
        "this is tracked as an open item."
    )

    # 7. Other requirements
    doc.add_heading("7. Other Requirements", level=1)
    doc.add_heading("7.1 Legal / compliance", level=2)
    doc.add_paragraph(
        "The system stores sensitive PII (PAN numbers, bank account details) "
        "with no field-level encryption beyond the database's own, and no "
        "data-retention or erasure policy is implemented. Statutory payroll "
        "compliance — EPFO/ESIC filings, TDS computation and remittance — is "
        "explicitly out of scope; PF and PT fields are user-entered "
        "deduction amounts only, not computed against statutory slabs."
    )
    doc.add_heading("7.2 Internationalization", level=2)
    doc.add_paragraph(
        "Currency formatting is fixed to Indian Rupees (en-IN locale) and is "
        "not currently configurable."
    )

    # Appendices
    doc.add_heading("Appendix A — Traceability to the Test Suite", level=1)
    doc.add_paragraph(
        "The automated Playwright suite (tests/e2e/) and the manual/automated "
        "test-case inventory (docs/SedPayroll-E2E-Test-Cases.xlsx) exercise "
        "the functional requirements in Section 3 by module: no-auth/ covers "
        "FR-AUTH, admin/ and employee/ cover FR-DASH through FR-PROF by role."
    )
    doc.add_heading("Appendix B — Glossary", level=1)
    req_table(doc, ["Term", "Meaning"], [
        ("CTC", "Cost to Company."),
        ("HRA", "House Rent Allowance."),
        ("PF / PT", "Provident Fund / Professional Tax."),
        ("LOP", "Loss of Pay."),
        ("RBAC", "Role-Based Access Control."),
        ("JWT", "JSON Web Token."),
    ], col_widths=[4, 12])

    add_page_number_footer(doc)
    return doc


def main():
    DOCS_DIR.mkdir(parents=True, exist_ok=True)

    brd = build_brd()
    brd_path = DOCS_DIR / "SedPayroll-BRD.docx"
    brd.save(brd_path)
    print(f"Wrote {brd_path.relative_to(ROOT)}")

    prd = build_prd()
    prd_path = DOCS_DIR / "SedPayroll-PRD.docx"
    prd.save(prd_path)
    print(f"Wrote {prd_path.relative_to(ROOT)}")

    srs = build_srs()
    srs_path = DOCS_DIR / "SedPayroll-SRS.docx"
    srs.save(srs_path)
    print(f"Wrote {srs_path.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
