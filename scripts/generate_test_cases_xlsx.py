#!/usr/bin/env python3
"""
Generates docs/SedPayroll-E2E-Test-Cases.xlsx — the manual test-case
inventory backing the Playwright suite under tests/e2e/.

Run: python3 scripts/generate_test_cases_xlsx.py
Requires: pip install openpyxl
"""
import subprocess
import sys
from pathlib import Path

try:
    import openpyxl  # noqa: F401
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", "openpyxl"])

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent.parent
OUT_PATH = ROOT / "docs" / "SedPayroll-E2E-Test-Cases.xlsx"

COLUMNS = [
    "ID", "Module", "Test Case", "Preconditions", "Test Steps",
    "Test Data", "Expected Result", "Role", "Priority", "Type",
    "Automated", "Automation Reference",
]

# Each row: (id, module, title, preconditions, steps, data, expected,
#            role, priority, type_, automated, reference)
ROWS = [
    # ---------------------------------------------------------------- AUTH
    ("AUTH-001", "Authentication", "Login page renders required fields and demo credential hints",
     "App is running; user is signed out.",
     "1. Navigate to /login.",
     "—",
     "Email and password fields, Sign in button, and both demo-credential hint lines are visible.",
     "Unauthenticated", "Medium", "Functional", "Yes", "tests/e2e/no-auth/login.spec.ts"),

    ("AUTH-002", "Authentication", "Admin logs in with valid credentials",
     "Seeded admin account exists (db:seed).",
     "1. Go to /login.\n2. Enter admin email + password.\n3. Click Sign in.",
     "neha.iyer@sedpayroll.com / Admin@123",
     "Redirected to \"/\"; sidebar visible; email + Administrator role shown; Employees/Payroll/Team Attendance nav items visible.",
     "Admin", "High", "Functional", "Yes", "tests/e2e/no-auth/login.spec.ts"),

    ("AUTH-003", "Authentication", "Employee logs in with valid credentials and sees a restricted nav",
     "Seeded employee account exists.",
     "1. Go to /login.\n2. Enter employee email + password.\n3. Click Sign in.",
     "aditi.sharma@sedpayroll.com / Employee@123",
     "Redirected to \"/\"; Employees/Payroll/Team Attendance nav items are absent.",
     "Employee", "High", "Functional", "Yes", "tests/e2e/no-auth/login.spec.ts"),

    ("AUTH-004", "Authentication", "Login fails with a wrong password",
     "Seeded account exists.",
     "1. Go to /login.\n2. Enter a valid email with an incorrect password.\n3. Click Sign in.",
     "neha.iyer@sedpayroll.com / wrong-password",
     "Stays on /login; shows \"Invalid email or password.\" No indication of whether the email itself was valid.",
     "Unauthenticated", "High", "Negative / Security", "Yes", "tests/e2e/no-auth/login.spec.ts"),

    ("AUTH-005", "Authentication", "Login fails for an email that doesn't exist",
     "—",
     "1. Go to /login.\n2. Enter an unregistered email + any password.\n3. Click Sign in.",
     "nobody@sedpayroll.com / wrong-password",
     "Stays on /login; shows the same generic \"Invalid email or password.\" message (no user-enumeration signal).",
     "Unauthenticated", "High", "Negative / Security", "Yes", "tests/e2e/no-auth/login.spec.ts"),

    ("AUTH-006", "Authentication", "Empty login form is blocked client-side",
     "—",
     "1. Go to /login.\n2. Click Sign in without filling either field.",
     "—",
     "Native \"required\" validation blocks submission; still on /login.",
     "Unauthenticated", "Low", "Negative", "Yes", "tests/e2e/no-auth/login.spec.ts"),

    ("AUTH-007", "Authentication", "Malformed email is rejected client-side",
     "—",
     "1. Go to /login.\n2. Enter \"not-an-email\" and any password.\n3. Click Sign in.",
     "not-an-email",
     "type=\"email\" validation blocks submission; still on /login.",
     "Unauthenticated", "Low", "Negative", "Yes", "tests/e2e/no-auth/login.spec.ts"),

    ("AUTH-008", "Authentication", "Password field masks input",
     "—",
     "1. Go to /login.\n2. Type into the password field.",
     "—",
     "Characters render as dots/bullets, not plain text.",
     "Unauthenticated", "Low", "UI", "No", "Manual — visual check"),

    ("AUTH-009", "Authentication", "An already-authenticated user visiting /login is redirected to the dashboard",
     "Signed in as any role.",
     "1. While signed in, navigate directly to /login.",
     "—",
     "Immediately redirected to \"/\" — never shown the login form again.",
     "Both", "Medium", "Functional", "No", "Manual — covered by lib/auth.config.ts's authorized() callback; add an automated check alongside global.setup.ts if this becomes flaky"),

    ("AUTH-010", "Authentication", "Signing out clears the session and returns to /login",
     "Signed in.",
     "1. Click \"Sign out\" in the app shell.",
     "—",
     "Redirected to /login; session cookie cleared.",
     "Both", "High", "Functional", "Yes", "tests/e2e/employee/dashboard.spec.ts"),

    ("AUTH-011", "Authentication", "Browser back-navigation after sign-out does not restore the session",
     "Signed in, then signed out.",
     "1. Sign out.\n2. Press the browser Back button.",
     "—",
     "Protected page re-checks the session server-side and redirects to /login rather than showing cached content.",
     "Both", "Medium", "Security", "No", "Manual — exploratory"),

    ("AUTH-012", "Authentication", "Injection-style strings in the email field are handled safely",
     "—",
     "1. Go to /login.\n2. Enter a string such as `' OR '1'='1` in the email field.\n3. Submit.",
     "' OR '1'='1",
     "Treated as an ordinary (invalid) email; generic \"Invalid email or password.\" error; no server error, no auth bypass (Drizzle's parameterized queries are expected to prevent injection).",
     "Unauthenticated", "Medium", "Security", "No", "Manual — exploratory security check"),

    # ---------------------------------------------------------------- RBAC
    ("RBAC-001", "Route protection", "Every portal route redirects an unauthenticated visitor to /login",
     "Signed out.",
     "1. Navigate directly to each of: /, /attendance, /payslips, /profile, /employees, /employees/new, /team-attendance, /payroll.",
     "—",
     "Every route redirects to /login (proxy.ts middleware).",
     "Unauthenticated", "High", "Security / RBAC", "Yes", "tests/e2e/no-auth/login.spec.ts"),

    ("RBAC-002", "Route protection", "Admin-only routes redirect a signed-in employee to the dashboard",
     "Signed in as an employee.",
     "1. Navigate directly to each of: /employees, /employees/new, /employees/1, /team-attendance, /payroll, /payroll/1.",
     "—",
     "Every route redirects to \"/\" (not /login — the user is authenticated, just not authorized).",
     "Employee", "High", "Security / RBAC", "Yes", "tests/e2e/employee/rbac.spec.ts"),

    ("RBAC-003", "Route protection", "Admin-only nav links are not rendered for an employee",
     "Signed in as an employee.",
     "1. Inspect the sidebar and mobile nav.",
     "—",
     "Employees / Team Attendance / Payroll links are absent from the DOM (not merely hidden).",
     "Employee", "Medium", "UI / RBAC", "Yes", "tests/e2e/employee/dashboard.spec.ts"),

    ("RBAC-004", "Route protection", "An admin retains access to both self-service and admin sections",
     "Signed in as admin.",
     "1. Inspect the sidebar.",
     "—",
     "All 7 nav items are present and each corresponding route loads.",
     "Admin", "High", "Functional / RBAC", "Yes", "tests/e2e/admin/dashboard.spec.ts"),

    ("RBAC-005", "Route protection", "An admin-only Server Action rejects a non-admin caller even if the UI is bypassed",
     "Signed in as an employee; a tool to submit a raw POST to a Server Action endpoint (e.g. browser devtools or curl with the session cookie).",
     "1. Attempt to invoke an admin-only action (e.g. runPayrollAction) directly, bypassing the UI.",
     "—",
     "`requireAdmin()` (lib/session.ts) rejects the call server-side (redirects/throws) — defense in depth beyond the proxy matcher, since Server Actions aren't covered by it.",
     "Employee", "High", "Security", "No", "Manual — exploratory, requires direct action invocation"),

    # ------------------------------------------------------------ DASHBOARD
    ("DASH-001", "Dashboard", "Employee dashboard shows self-service widgets only",
     "Signed in as employee.",
     "1. Go to \"/\".",
     "—",
     "Clock widget and latest-payslip card are visible; organization overview is absent.",
     "Employee", "Medium", "Functional", "Yes", "tests/e2e/employee/dashboard.spec.ts"),

    ("DASH-002", "Dashboard", "Admin dashboard additionally shows the organization overview",
     "Signed in as admin.",
     "1. Go to \"/\".",
     "—",
     "Clock widget, latest-payslip card, and a 4-stat organization overview (Active employees, Present today, Payroll runs, Inactive employees) are all visible.",
     "Admin", "Medium", "Functional", "Yes", "tests/e2e/admin/dashboard.spec.ts"),

    ("DASH-003", "Dashboard", "Organization overview stat cards link to their section",
     "Signed in as admin.",
     "1. Click the \"Active employees\" stat card.",
     "—",
     "Navigates to /employees.",
     "Admin", "Low", "Functional", "Yes", "tests/e2e/admin/dashboard.spec.ts"),

    ("DASH-004", "Dashboard", "Sidebar reflects the signed-in user's email and role label",
     "Signed in as either role.",
     "1. Inspect the sidebar footer.",
     "—",
     "Shows the correct email and \"Administrator\"/\"Employee\" label matching the account.",
     "Both", "Low", "UI", "Yes", "tests/e2e/admin/dashboard.spec.ts, tests/e2e/employee/dashboard.spec.ts"),

    ("DASH-005", "Dashboard", "Latest-payslip card reflects the most recent payslip, or an empty state pre-payroll",
     "Employee with/without a generated payslip.",
     "1. Go to \"/\" as an employee who has a payslip.\n2. Repeat for one who has none.",
     "—",
     "Shows net pay when a payslip exists; shows \"No payslips yet\" copy otherwise.",
     "Employee", "Low", "Functional / Data", "No", "Manual — depends on seeded/generated payroll state"),

    ("DASH-006", "Dashboard", "Mobile viewport shows the mobile nav instead of the desktop sidebar",
     "Any signed-in user.",
     "1. Resize the viewport to a phone width (e.g. 375px).\n2. Observe navigation.",
     "—",
     "Desktop sidebar is hidden; horizontal mobile nav bar is shown and functional.",
     "Both", "Low", "UI / Responsive", "No", "Manual — responsive/visual exploratory"),

    # ------------------------------------------------------------ ATTENDANCE
    ("ATT-001", "My Attendance", "Attendance history table renders",
     "Signed in as employee.",
     "1. Go to /attendance.",
     "—",
     "Table renders (rows if history exists, otherwise \"No attendance records yet.\").",
     "Employee", "Medium", "Functional", "Yes", "tests/e2e/employee/attendance.spec.ts"),

    ("ATT-002", "My Attendance", "Clocking in records today's session and updates the status text",
     "Employee has not clocked in today.",
     "1. Go to \"/\".\n2. Click \"Clock in\".",
     "—",
     "Status text changes to \"Clocked in at HH:MM\"; a new attendance row is created for today.",
     "Employee", "High", "Functional", "Yes", "tests/e2e/employee/attendance.spec.ts"),

    ("ATT-003", "My Attendance", "\"Clock in\" is disabled once already clocked in today",
     "Employee has clocked in today.",
     "1. Go to \"/\".\n2. Observe the Clock in button.",
     "—",
     "Button is disabled — cannot submit a second clock-in via the UI.",
     "Employee", "Medium", "Functional / Negative", "Yes", "tests/e2e/employee/attendance.spec.ts"),

    ("ATT-004", "My Attendance", "\"Clock out\" is disabled until clocked in",
     "Employee has not clocked in today.",
     "1. Go to \"/\".\n2. Observe the Clock out button.",
     "—",
     "Button is disabled.",
     "Employee", "Medium", "Functional / Negative", "Yes", "tests/e2e/employee/attendance.spec.ts"),

    ("ATT-005", "My Attendance", "Clocking out updates status and disables further clock-out",
     "Employee has clocked in today, not yet clocked out.",
     "1. Click \"Clock out\".",
     "—",
     "Status text changes to \"Worked HH:MM – HH:MM\"; Clock out button becomes disabled.",
     "Employee", "High", "Functional", "Yes", "tests/e2e/employee/attendance.spec.ts"),

    ("ATT-006", "My Attendance", "Clocking out under 5 worked hours marks the day half_day",
     "Employee clocked in less than 5 hours ago.",
     "1. Clock in.\n2. (Wait / simulate <5 hours elapsed.)\n3. Clock out.",
     "—",
     "Attendance status is \"half_day\" for today.",
     "Employee", "Medium", "Functional / Business rule", "No", "Manual — requires time manipulation to exercise reliably"),

    ("ATT-007", "My Attendance", "Clocking out at 5+ worked hours marks the day present",
     "Employee clocked in 5+ hours ago.",
     "1. Clock in.\n2. (Wait / simulate ≥5 hours elapsed.)\n3. Clock out.",
     "—",
     "Attendance status is \"present\" for today.",
     "Employee", "Medium", "Functional / Business rule", "No", "Manual — requires time manipulation to exercise reliably"),

    ("ATT-008", "My Attendance", "Today's session appears in the attendance history",
     "Employee clocked in today.",
     "1. Clock in.\n2. Go to /attendance.",
     "—",
     "A row for today's date is present.",
     "Employee", "Medium", "Functional", "Yes", "tests/e2e/employee/attendance.spec.ts"),

    ("ATT-009", "My Attendance", "A second clock-in on the same day is rejected server-side",
     "Employee already clocked in today; UI button bypassed (e.g. direct action call).",
     "1. Invoke clockInAction() a second time for the same day.",
     "—",
     "Returns \"You have already clocked in today.\" and does not insert a duplicate row (defense in depth beyond the disabled button).",
     "Employee", "Medium", "Negative / Security", "No", "Manual — exploratory, requires bypassing the disabled button"),

    ("ATT-010", "My Attendance", "Clocking out without having clocked in is rejected server-side",
     "Employee has not clocked in today; UI button bypassed.",
     "1. Invoke clockOutAction() without a prior clock-in.",
     "—",
     "Returns \"You need to clock in first.\"",
     "Employee", "Medium", "Negative", "No", "Manual — exploratory, requires bypassing the disabled button"),

    # ------------------------------------------------------- TEAM ATTENDANCE
    ("TATT-001", "Team Attendance", "Team attendance defaults to today, listing all active employees",
     "Signed in as admin.",
     "1. Go to /team-attendance.",
     "—",
     "Table title shows today's date; one row per active employee.",
     "Admin", "Medium", "Functional", "Yes", "tests/e2e/admin/team-attendance.spec.ts"),

    ("TATT-002", "Team Attendance", "Date picker filters attendance to a chosen past date",
     "Signed in as admin.",
     "1. Open the date picker.\n2. Select yesterday.\n3. Click View.",
     "—",
     "URL updates to ?date=YYYY-MM-DD; table title and rows reflect the selected date.",
     "Admin", "Medium", "Functional", "Yes", "tests/e2e/admin/team-attendance.spec.ts"),

    ("TATT-003", "Team Attendance", "Inactive employees are excluded from the team attendance list",
     "At least one inactive employee exists.",
     "1. Deactivate an employee.\n2. Go to /team-attendance.",
     "—",
     "The deactivated employee's row is absent regardless of selected date.",
     "Admin", "Low", "Functional / Data", "No", "Manual — data-dependent; not isolated as its own automated case, though covered incidentally by employees.spec.ts's deactivate flow"),

    ("TATT-004", "Team Attendance", "Employees with no attendance record for the selected date show placeholders",
     "An active employee has no attendance row for the selected date.",
     "1. Go to /team-attendance for a date with a gap.",
     "—",
     "Clock in / Clock out cells render \"—\" rather than blank or an error.",
     "Admin", "Low", "UI", "No", "Manual — visual check"),

    # ------------------------------------------------------- EMPLOYEE MGMT
    ("EMP-001", "Employee Management", "Employee list shows active employees by default",
     "Signed in as admin.",
     "1. Go to /employees.",
     "—",
     "Table lists active employees; Active tab is selected.",
     "Admin", "Medium", "Functional", "Yes", "tests/e2e/admin/employees.spec.ts"),

    ("EMP-002", "Employee Management", "Active/Inactive tabs filter the roster",
     "Signed in as admin.",
     "1. Click the Inactive tab.\n2. Click the Active tab.",
     "—",
     "URL and table contents update to match the selected status filter.",
     "Admin", "Medium", "Functional", "Yes", "tests/e2e/admin/employees.spec.ts"),

    ("EMP-003", "Employee Management", "Create-employee form rejects an empty submission",
     "Signed in as admin.",
     "1. Go to /employees/new.\n2. Click Create employee without filling anything.",
     "—",
     "Generic \"Please fix the highlighted fields.\" banner; per-field errors under First name, Last name, Email, Date of joining; no employee created.",
     "Admin", "High", "Negative", "Yes", "tests/e2e/admin/employees.spec.ts"),

    ("EMP-004", "Employee Management", "Create-employee form rejects a duplicate email",
     "An employee with the target email already exists.",
     "1. Fill the form with an existing employee's email plus valid required fields.\n2. Submit.",
     "aditi.sharma@sedpayroll.com",
     "\"An account with this email already exists.\" banner and field error; no duplicate account created.",
     "Admin", "High", "Negative / Data integrity", "Yes", "tests/e2e/admin/employees.spec.ts"),

    ("EMP-005", "Employee Management", "Create a new employee (employee role)",
     "Signed in as admin.",
     "1. Fill required fields (name, email, gender, department, designation, date of joining, basic salary).\n2. Submit.",
     "Unique generated email",
     "Redirects to /employees/{id}?created=1; shows \"Employee@123\" as the temp password; employee is auto-assigned an employee code and appears in the active roster.",
     "Admin", "High", "Functional", "Yes", "tests/e2e/admin/employees.spec.ts"),

    ("EMP-006", "Employee Management", "Create a new employee with \"Grant admin access\" checked",
     "Signed in as admin.",
     "1. Fill required fields.\n2. Check \"Grant admin access\".\n3. Submit.",
     "Unique generated email",
     "Shows \"Admin@123\" as the temp password; the employee's role badge reads Administrator.",
     "Admin", "High", "Functional", "Yes", "tests/e2e/admin/employees.spec.ts"),

    ("EMP-007", "Employee Management", "Editing an employee persists field changes",
     "An employee exists.",
     "1. Open an employee's edit page.\n2. Change a field (e.g. phone).\n3. Submit.",
     "9876543210",
     "Form reloads with the new value persisted (visible on reload).",
     "Admin", "Medium", "Functional", "Yes", "tests/e2e/admin/employees.spec.ts"),

    ("EMP-008", "Employee Management", "Deactivating an employee moves them to the Inactive tab",
     "An active employee exists.",
     "1. Click Deactivate on their row.",
     "—",
     "Row disappears from the Active tab; appears under Inactive; date of leaving is set.",
     "Admin", "High", "Functional", "Yes", "tests/e2e/admin/employees.spec.ts"),

    ("EMP-009", "Employee Management", "Reactivating an employee moves them back to the Active tab",
     "An inactive employee exists.",
     "1. From the Inactive tab, click Reactivate on their row.",
     "—",
     "Row disappears from Inactive; reappears under Active.",
     "Admin", "Medium", "Functional", "Yes", "tests/e2e/admin/employees.spec.ts"),

    ("EMP-010", "Employee Management", "An admin cannot remove their own admin access",
     "Signed in as admin, viewing their own edit page.",
     "1. Open your own employee edit page.\n2. Observe the admin-access toggle.",
     "—",
     "Toggle button is disabled with the note \"You can't change your own admin access.\"",
     "Admin", "High", "Negative / Security", "Yes", "tests/e2e/admin/employees.spec.ts"),

    ("EMP-011", "Employee Management", "An admin can grant/revoke another admin's access",
     "A second admin account exists.",
     "1. Open a different admin's edit page.\n2. Click \"Remove admin access\", confirm badge changes.\n3. Click \"Grant admin access\" to restore.",
     "—",
     "Role badge toggles between Administrator and Employee correctly each time.",
     "Admin", "High", "Functional", "Yes", "tests/e2e/admin/employees.spec.ts"),

    ("EMP-012", "Employee Management", "The last remaining admin cannot be demoted",
     "Exactly one admin exists in the system.",
     "1. Attempt to remove admin access from the sole remaining admin.",
     "—",
     "Rejected with \"There must be at least one admin.\"",
     "Admin", "High", "Negative / Security", "No",
     "Manual — this guard in lib/actions/employees.ts is checked after the self-demotion guard, and is only reachable when the acting admin and the sole admin are different accounts, which cannot occur through the UI in this single-admin scenario; verify via direct DB/action-level testing"),

    ("EMP-013", "Employee Management", "Numeric salary fields reject negative values",
     "Signed in as admin, on the create-employee form.",
     "1. Enter a negative number in Basic salary (e.g. -1000).\n2. Submit.",
     "-1000",
     "Validation rejects the value (zod `.min(0)`); form does not submit successfully.",
     "Admin", "Medium", "Negative / Validation", "No", "Manual — quick to verify, not yet automated"),

    ("EMP-014", "Employee Management", "Employee code is generated sequentially",
     "Multiple employees created over time.",
     "1. Create several employees in sequence.\n2. Compare their employee codes.",
     "—",
     "Codes follow EMP001, EMP002, … with no gaps or collisions.",
     "Admin", "Low", "Functional / Data", "No", "Manual — data/sequence verification"),

    ("EMP-015", "Employee Management", "Salary structure CTC is computed correctly",
     "A new employee/salary structure is created.",
     "1. Create an employee with known basic/HRA/conveyance/medical/special values.\n2. Inspect the stored salary structure (e.g. via Drizzle Studio).",
     "e.g. basic 50000, hra 20000, others 0",
     "CTC = (basic + hra + conveyance + medical + special) × 12.",
     "Admin", "Low", "Functional / Calculation", "No", "Manual — calculation verification, requires DB inspection"),

    ("EMP-016", "Employee Management", "Editing an employee's email to one already used elsewhere is rejected",
     "Two employees exist.",
     "1. Edit employee A's email to match employee B's email.\n2. Submit.",
     "—",
     "\"An account with this email already exists.\" banner; email not changed.",
     "Admin", "Medium", "Negative / Data integrity", "No", "Manual — symmetrical to EMP-004's create-time check, not yet automated for edit"),

    ("EMP-017", "Employee Management", "Leaving Gender / Department / Designation unselected blocks submission with no visible reason",
     "Signed in as admin, on the create-employee form.",
     "1. Fill only the fields with visible required-asterisk styling (first/last name, email, date of joining, basic).\n2. Leave Gender, Department, and Designation unselected.\n3. Submit.",
     "—",
     "KNOWN ISSUE: submission fails with only the generic \"Please fix the highlighted fields.\" banner — no per-field error appears under Gender, Department, or Designation, even though all three are effectively required in practice (an unselected shadcn Select submits `\"\"`, which fails the zod schema's enum/positive-int check rather than being treated as omitted). A user following the visible field indicators has no way to tell why the form failed. Recommend either truly making these optional in `lib/validation.ts`, or wiring `error={fieldError(...)}` into their `Field` wrappers in `employee-form.tsx` like every other field.",
     "Admin", "High", "Bug / Validation", "No", "Manual — logged from building this suite; see tests/support/helpers.ts's fillRequiredEmployeeFields() comment for the workaround the automated suite uses"),

    # -------------------------------------------------------------- PAYROLL
    ("PAY-001", "Payroll", "Payroll page shows the run form and historical runs",
     "Signed in as admin.",
     "1. Go to /payroll.",
     "—",
     "Run-payroll form and a table of past runs are both visible.",
     "Admin", "Medium", "Functional", "Yes", "tests/e2e/admin/payroll.spec.ts"),

    ("PAY-002", "Payroll", "Running payroll for a new month succeeds",
     "No payroll run exists yet for the chosen month/year.",
     "1. Select month and year.\n2. Click \"Run payroll\".",
     "Current month/year",
     "Redirects to /payroll/{id}; status \"processed\"; a payslip row exists for every active employee.",
     "Admin", "High", "Functional", "Yes", "tests/e2e/admin/payroll.spec.ts"),

    ("PAY-003", "Payroll", "Running payroll for an already-processed month is rejected",
     "A payroll run already exists for the chosen month/year (seed always processes last month).",
     "1. Select last month + this year.\n2. Click \"Run payroll\".",
     "Last month / this year",
     "\"Payroll for this month has already been run.\" error; no duplicate run created; stays on /payroll.",
     "Admin", "High", "Negative / Data integrity", "Yes", "tests/e2e/admin/payroll.spec.ts"),

    ("PAY-004", "Payroll", "Payroll run detail lists per-employee payslips and totals",
     "At least one payroll run exists.",
     "1. Click \"View payslips\" on a run.",
     "—",
     "Payslips table lists every employee paid in that run with days present, gross earnings, deductions, and net pay.",
     "Admin", "High", "Functional", "Yes", "tests/e2e/admin/payroll.spec.ts"),

    ("PAY-005", "Payroll", "Marking a processed run as paid updates its status",
     "A run with status \"processed\" exists.",
     "1. Open the run.\n2. Click \"Mark as paid\".",
     "—",
     "Status badge changes to \"paid\"; the \"Mark as paid\" button disappears.",
     "Admin", "High", "Functional", "Yes", "tests/e2e/admin/payroll.spec.ts"),

    ("PAY-006", "Payroll", "\"Mark as paid\" is only available for processed runs",
     "Runs in different statuses (processed vs paid) exist.",
     "1. Open a \"paid\" run.\n2. Observe available actions.",
     "—",
     "No \"Mark as paid\" button is shown once a run is already paid.",
     "Admin", "Medium", "Functional", "Yes", "tests/e2e/admin/payroll.spec.ts (asserted as part of PAY-005)"),

    ("PAY-007", "Payroll", "Visiting a non-existent payroll run returns 404",
     "Signed in as admin.",
     "1. Navigate to /payroll/999999.",
     "—",
     "HTTP 404 / Next.js not-found page.",
     "Admin", "Medium", "Negative", "Yes", "tests/e2e/admin/payroll.spec.ts"),

    ("PAY-008", "Payroll", "Payslip amounts are pro-rated for partial attendance",
     "An employee has fewer days present than days in the month.",
     "1. Run payroll for a month with a known partial-attendance employee.\n2. Inspect their payslip.",
     "—",
     "Gross earnings / net pay reflect proportional pay for days present (per lib/payroll.ts's computePayslipForEmployee).",
     "Admin", "Medium", "Functional / Calculation", "No", "Manual — calculation verification against known attendance data"),

    ("PAY-009", "Payroll", "An employee marked on_leave for the full month is paid per the on-leave business rule",
     "An employee has on_leave attendance status for the relevant days.",
     "1. Run payroll for that month.\n2. Inspect their payslip.",
     "—",
     "on_leave days are treated as paid per lib/payroll.ts's rule (not deducted as loss of pay).",
     "Admin", "Low", "Functional / Business rule", "No", "Manual — calculation verification"),

    ("PAY-010", "Payroll", "Payroll run's total net pay equals the sum of its payslips",
     "A payroll run with multiple payslips exists.",
     "1. Open the run detail page.\n2. Sum the \"Net pay\" column.\n3. Compare to the header's \"Total net pay\".",
     "—",
     "The two values match exactly.",
     "Admin", "Low", "Functional / Calculation", "No", "Manual — arithmetic cross-check"),

    # ------------------------------------------------------------- PAYSLIPS
    ("PSLIP-001", "My Payslips", "Payslip list shows the employee's own payslips or an empty state",
     "Signed in as employee.",
     "1. Go to /payslips.",
     "—",
     "Table lists the employee's payslips (month, gross, net) or shows \"No payslips yet.\"",
     "Employee", "Medium", "Functional", "Yes", "tests/e2e/employee/payslips.spec.ts"),

    ("PSLIP-002", "My Payslips", "Opening a payslip shows the full earnings/deductions breakdown",
     "The employee has at least one payslip.",
     "1. Click \"View\" on a payslip row.",
     "—",
     "Detail page shows Earnings and Deductions sections and a Net pay total consistent with the list.",
     "Employee", "High", "Functional", "Yes", "tests/e2e/employee/payslips.spec.ts"),

    ("PSLIP-003", "My Payslips", "An employee cannot view another employee's payslip by guessing its id",
     "Signed in as employee; other employees have payslips with nearby ids.",
     "1. Navigate to /payslips/{id} for ids that don't belong to this employee.",
     "IDs 1–20",
     "Every id that isn't this employee's own payslip returns 404 — never another employee's data.",
     "Employee", "High", "Security", "Yes", "tests/e2e/employee/payslips.spec.ts"),

    ("PSLIP-004", "My Payslips", "Numbers on an employee's own payslip match what the admin sees for the same payslip",
     "A payroll run has been processed.",
     "1. As admin, open the run detail and note one employee's figures.\n2. As that employee, open the equivalent payslip.",
     "—",
     "Gross earnings, deductions, and net pay match exactly between both views.",
     "Both", "Low", "Functional / Cross-check", "No", "Manual — cross-role consistency check"),

    ("PSLIP-005", "My Payslips", "Net pay equals gross earnings minus total deductions",
     "An employee has a payslip.",
     "1. Open the payslip detail page.\n2. Compute gross − deductions by hand.",
     "—",
     "Result matches the displayed Net pay exactly.",
     "Employee", "Low", "Functional / Calculation", "No", "Manual — arithmetic cross-check"),

    # -------------------------------------------------------------- PROFILE
    ("PROF-001", "Profile", "Profile page shows correct read-only employee details",
     "Signed in as employee.",
     "1. Go to /profile.",
     "—",
     "Shows correct name, email, department, designation, date of joining/birth, and address for the signed-in employee.",
     "Employee", "Medium", "Functional", "Yes", "tests/e2e/employee/profile.spec.ts"),

    ("PROF-002", "Profile", "Profile page has no editable fields",
     "Signed in as employee.",
     "1. Go to /profile.\n2. Inspect for input/textarea/select elements.",
     "—",
     "No editable form controls are present (contact-HR copy is shown instead).",
     "Employee", "Low", "UI", "Yes", "tests/e2e/employee/profile.spec.ts"),

    # ---------------------------------------------------------- NON-FUNCT'L
    ("NFR-001", "Non-functional", "App renders correctly across major browser engines",
     "Playwright browsers installed (chromium, firefox, webkit).",
     "1. Run the suite with --project=chromium/firefox/webkit (add matching projects to playwright.config.ts).",
     "—",
     "Core flows (login, dashboard, employee CRUD, payroll) pass on all three engines.",
     "Both", "Low", "Compatibility", "No", "Manual/CI extension — config currently ships Chromium-based projects only"),

    ("NFR-002", "Non-functional", "Responsive layout holds up on a mobile viewport",
     "—",
     "1. Resize to common phone/tablet widths.\n2. Exercise the main flows.",
     "—",
     "No horizontal overflow; mobile nav is usable; tables scroll within their own container.",
     "Both", "Low", "UI / Responsive", "No", "Manual — visual exploratory"),

    ("NFR-003", "Non-functional", "Keyboard-only navigation reaches all primary actions",
     "—",
     "1. Navigate the app using only Tab/Shift+Tab/Enter.\n2. Use the \"Skip to main content\" link.",
     "—",
     "Focus order is logical; all nav items and form controls are reachable; skip link jumps to <main>.",
     "Both", "Medium", "Accessibility", "No", "Manual — accessibility audit"),

    ("NFR-004", "Non-functional", "No unexpected console errors on primary pages",
     "—",
     "1. Open devtools console.\n2. Visit every primary route as both roles.",
     "—",
     "No uncaught errors/warnings beyond expected framework logs.",
     "Both", "Low", "Quality", "No", "Manual — exploratory"),

    ("NFR-005", "Non-functional", "Submit buttons show a loading state and prevent double-submission",
     "—",
     "1. Submit a form (e.g. login, create employee).\n2. Observe the button during the pending request.",
     "—",
     "Button shows a spinner and is disabled until the action resolves — a rapid double-click only submits once.",
     "Both", "Medium", "UI / Robustness", "No", "Manual — exploratory, timing-sensitive"),

    ("NFR-006", "Non-functional", "Currency and date formatting is consistent across the app",
     "—",
     "1. Compare currency/date rendering on the dashboard, payslip detail, payroll run detail, and employee list.",
     "—",
     "All amounts render in the same locale/currency format (see lib/format.ts); dates render consistently.",
     "Both", "Low", "UI / Consistency", "No", "Manual — visual check"),
]


def build_workbook() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Test Cases"

    header_fill = PatternFill("solid", fgColor="1F2937")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    wrap = Alignment(wrap_text=True, vertical="top", horizontal="left")
    center = Alignment(wrap_text=True, vertical="top", horizontal="center")
    thin = Side(style="thin", color="D1D5DB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = border
    ws.row_dimensions[1].height = 22

    priority_fill = {
        "High": PatternFill("solid", fgColor="FEE2E2"),
        "Medium": PatternFill("solid", fgColor="FEF3C7"),
        "Low": PatternFill("solid", fgColor="E5E7EB"),
    }
    automated_fill = {
        "Yes": PatternFill("solid", fgColor="D1FAE5"),
        "No": PatternFill("solid", fgColor="F3F4F6"),
    }

    centered_cols = {1, 8, 9, 10, 11}  # ID, Role, Priority, Type, Automated

    for row in ROWS:
        ws.append(row)
        r = ws.max_row
        for c in range(1, len(COLUMNS) + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = border
            cell.alignment = center if c in centered_cols else wrap
        ws.cell(row=r, column=9).fill = priority_fill.get(row[8], PatternFill())
        ws.cell(row=r, column=11).fill = automated_fill.get(row[10], PatternFill())

    widths = [10, 18, 34, 26, 34, 20, 40, 12, 10, 16, 11, 34]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    last_row = ws.max_row
    last_col = get_column_letter(len(COLUMNS))
    table = Table(displayName="TestCases", ref=f"A1:{last_col}{last_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False
    )
    ws.add_table(table)

    # ---- Summary sheet ----
    summary = wb.create_sheet("Summary")
    summary.append(["SedPayroll E2E Test Case Inventory"])
    summary["A1"].font = Font(bold=True, size=14)
    summary.append([])

    total = len(ROWS)
    automated = sum(1 for r in ROWS if r[10] == "Yes")
    manual = total - automated
    summary.append(["Total test cases", total])
    summary.append(["Automated (Playwright)", automated])
    summary.append(["Manual / exploratory", manual])
    summary.append([])

    summary.append(["By module", "Count"])
    modules = {}
    for r in ROWS:
        modules[r[1]] = modules.get(r[1], 0) + 1
    for m, c in modules.items():
        summary.append([m, c])
    summary.append([])

    summary.append(["Automated spec files (tests/e2e/)"])
    for path in sorted({
        "tests/e2e/no-auth/login.spec.ts",
        "tests/e2e/admin/dashboard.spec.ts",
        "tests/e2e/admin/employees.spec.ts",
        "tests/e2e/admin/payroll.spec.ts",
        "tests/e2e/admin/team-attendance.spec.ts",
        "tests/e2e/employee/dashboard.spec.ts",
        "tests/e2e/employee/attendance.spec.ts",
        "tests/e2e/employee/payslips.spec.ts",
        "tests/e2e/employee/profile.spec.ts",
        "tests/e2e/employee/rbac.spec.ts",
    }):
        summary.append([path])

    summary.column_dimensions["A"].width = 40
    summary.column_dimensions["B"].width = 14
    for cell in summary["A"]:
        cell.alignment = Alignment(horizontal="left")

    return wb


def main():
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb = build_workbook()
    wb.save(OUT_PATH)
    print(f"Wrote {len(ROWS)} test cases to {OUT_PATH.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
